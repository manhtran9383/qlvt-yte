"""
Import dữ liệu từ file Excel V8-ADJ vào qlvt.db.
File này có sheet 'in out' với 3,537 dòng lịch sử giao dịch đầy đủ từ 2020.
"""
import os, sys, sqlite3
from datetime import datetime
from openpyxl import load_workbook

EXCEL = r"C:\Users\manht\AppData\Local\hermes\attachments\2026.10.07 Quan ly vat tu V8 - ADJ.xlsm"
DB = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "qlvt.db")

def norm(s):
    return "" if s is None else str(s).strip().lower()

def to_float(v, d=0):
    try: return float(v)
    except: return d

def xdate(v):
    if v is None: return None
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try: return datetime.strptime(s, f).strftime("%Y-%m-%d")
        except: pass
    return None

print(f"Đọc: {EXCEL}")
wb = load_workbook(EXCEL, data_only=True, read_only=True)

conn = sqlite3.connect(DB)
conn.execute("PRAGMA journal_mode=WAL")
conn.execute("PRAGMA foreign_keys=OFF")

# === XÓA DỮ LIỆU CŨ (giữ users) ===
conn.execute("DELETE FROM phieu_giaodich")
conn.execute("DELETE FROM hop_dong")
conn.execute("DELETE FROM bang_gia")
conn.execute("DELETE FROM dm_khachhang")
conn.execute("DELETE FROM dm_vattu")
conn.commit()
print("Đã xóa dữ liệu cũ")

# === PRODUCT: header R2, data R3+ ===
ws = wb["Product"]
n = 0
for row in list(ws.iter_rows(values_only=True))[2:]:
    if not row or not row[0]: continue
    ma = str(row[0]).strip()
    conn.execute("INSERT OR REPLACE INTO dm_vattu VALUES (?,?,?,?,?,?,?,?,?)",
        (ma, str(row[1] or "").strip(), str(row[6] or "").strip() if len(row)>6 else "",
         str(row[2] or "").strip(), str(row[3] or "").strip(), "", "",
         to_float(row[10] if len(row)>10 else 0), 0))
    n += 1
print(f"Product: {n} dòng")

# === CUSTOMER: header R2, data R3+ ===
ws = wb["Customer"]
n = 0
for row in list(ws.iter_rows(values_only=True))[2:]:
    if not row or not row[1]: continue
    conn.execute("INSERT OR REPLACE INTO dm_khachhang VALUES (?,?,?,?,?)",
        (str(row[1]).strip(), str(row[0] or "").strip(),
         str(row[2] or "").strip() if len(row)>2 else "",
         str(row[3] or "").strip() if len(row)>3 else "",
         str(row[4] or "").strip() if len(row)>4 else ""))
    n += 1
print(f"Customer: {n} dòng")

ten_to_ma = {norm(r[1]): r[0] for r in conn.execute("SELECT ma_kh,ten_kh FROM dm_khachhang")}
ma_vt_set = {r[0] for r in conn.execute("SELECT ma_vt FROM dm_vattu")}

# === IN OUT: lịch sử giao dịch đầy đủ ===
# Cột: 0=Thời gian, 1=Mã VT, 2=Tên VT, 3=ĐVT, 4=Hạn dùng, 5=Nhập, 6=Xuất, 7=Ký gửi, 8=Xuất kho ký gửi, 9=Đơn giá, 10=Thành tiền
ws = wb["in out"]
n_nhap = n_xuat = n_kygui = n_xuat_kygui = 0
for row in list(ws.iter_rows(min_row=3, values_only=True)):
    if not row or (row[0] is None and row[1] is None): continue
    ma = str(row[1] or "").strip()
    if not ma or ma not in ma_vt_set: continue
    ngay = xdate(row[0]) or "2026-01-01"
    hsd = xdate(row[4] if len(row)>4 else None)
    sl_nhap = to_float(row[5] if len(row)>5 else 0)
    sl_xuat = to_float(row[6] if len(row)>6 else 0)
    sl_kygui = to_float(row[7] if len(row)>7 else 0)
    sl_xuat_kygui = to_float(row[8] if len(row)>8 else 0)
    gia = to_float(row[9] if len(row)>9 else 0)
    now = datetime.now().strftime("%Y-%m-%d")
    if sl_nhap > 0:
        conn.execute("INSERT INTO phieu_giaodich (so_phieu,loai_phieu,ngay_phieu,ma_vt,han_su_dung,so_luong,don_gia,ngay_lap) VALUES (?,?,?,?,?,?,?,?)",
            (f"IMP-{n_nhap+1:05d}", "NHAP", ngay, ma, hsd, sl_nhap, gia, now))
        n_nhap += 1
    if sl_xuat > 0:
        conn.execute("INSERT INTO phieu_giaodich (so_phieu,loai_phieu,ngay_phieu,ma_vt,han_su_dung,so_luong,don_gia,ngay_lap) VALUES (?,?,?,?,?,?,?,?)",
            (f"SAL-{n_xuat+1:05d}", "XUAT", ngay, ma, hsd, sl_xuat, gia, now))
        n_xuat += 1
    if sl_kygui > 0:
        conn.execute("INSERT INTO phieu_giaodich (so_phieu,loai_phieu,ngay_phieu,ma_vt,han_su_dung,so_luong,don_gia,ngay_lap) VALUES (?,?,?,?,?,?,?,?)",
            (f"KG-{n_kygui+1:05d}", "KYGUI", ngay, ma, hsd, sl_kygui, gia, now))
        n_kygui += 1
    if sl_xuat_kygui > 0:
        conn.execute("INSERT INTO phieu_giaodich (so_phieu,loai_phieu,ngay_phieu,ma_vt,han_su_dung,so_luong,don_gia,ngay_lap) VALUES (?,?,?,?,?,?,?,?)",
            (f"XKG-{n_xuat_kygui+1:05d}", "XUAT_KYGUI", ngay, ma, hsd, sl_xuat_kygui, gia, now))
        n_xuat_kygui += 1
print(f"In out: nhap={n_nhap}, xuat={n_xuat}, kygui={n_kygui}, xuat_kygui={n_xuat_kygui}")

# === PRDETAIL: bảng giá theo bệnh viện ===
# Header R3: cột 0=mã VT, cột 1+ = tên bệnh viện, data R4+
ws = wb["PrDetail"]
rows = list(ws.iter_rows(values_only=True))
header = rows[2] if len(rows)>2 else None
if header:
    khs = []
    for i, cell in enumerate(header[1:], 1):
        if cell and str(cell).strip():
            ten = str(cell).strip()
            ma_kh = next((k for k, v in ten_to_ma.items() if norm(v)==norm(ten)), None)
            if not ma_kh and ten in [r[1] for r in conn.execute("SELECT ma_kh,ten_kh FROM dm_khachhang")]:
                ma_kh = next((r[0] for r in conn.execute("SELECT ma_kh,ten_kh FROM dm_khachhang") if r[1]==ten), None)
            if not ma_kh:
                # Tìm theo tên chứa trong tên KH
                for k, v in ten_to_ma.items():
                    if norm(ten) in norm(v) or norm(v) in norm(ten):
                        ma_kh = k; break
            if ma_kh:
                khs.append((i, ma_kh))
    n = 0
    for row in rows[3:]:
        if not row or not row[0]: continue
        ma = str(row[0]).strip()
        for i, ma_kh in khs:
            if i < len(row) and row[i]:
                gia = to_float(row[i])
                if gia > 0:
                    conn.execute("INSERT OR REPLACE INTO bang_gia VALUES (?,?,?)", (ma, ma_kh, gia))
                    n += 1
    print(f"PrDetail (bang gia): {n} dong")

# === CONTRACTS: header R3, data R4+ ===
ws = wb["Contracts"]
n = 0
for row in list(ws.iter_rows(values_only=True))[3:]:
    if not row or not row[1]: continue
    kh = str(row[0] or "").strip()
    hd = str(row[1] or "").strip()
    ma = str(row[9] or "").strip() if len(row)>9 else ""
    ma_kh = ten_to_ma.get(norm(kh))
    if not ma_kh or not ma: continue
    conn.execute("INSERT OR REPLACE INTO hop_dong (so_hd,ma_kh,ma_vt,so_luong,don_gia,ngay_ky,ngay_het_han) VALUES (?,?,?,?,?,?,?)",
        (hd, ma_kh, ma, 0, 0, xdate(row[4] if len(row)>4 else None), xdate(row[6] if len(row)>6 else None)))
    n += 1
print(f"Contracts: {n}")

# === INVOICE IN ADVANCE: header R3, data R4+ ===
ws = wb["Invoice in advance"]
n = 0
for row in list(ws.iter_rows(values_only=True))[3:]:
    if not row or not row[4]: continue
    kh = str(row[0] or "").strip()
    ma = str(row[4] or "").strip()
    sl = to_float(row[7] if len(row)>7 else 0)
    ngay = xdate(row[3] if len(row)>3 else None) or "2026-01-01"
    ma_kh = ten_to_ma.get(norm(kh))
    if not ma_kh or not ma: continue
    conn.execute("INSERT INTO phieu_giaodich (so_phieu,loai_phieu,ngay_phieu,ma_vt,so_luong,don_gia,ma_kh,so_hop_dong,ngay_lap) VALUES (?,?,?,?,?,?,?,?,?)",
        (f"HDT-{n+1:05d}", "XUAT_HOADON_TRUOC", ngay, ma, sl, 0, ma_kh, str(row[1] or "").strip(), datetime.now().strftime("%Y-%m-%d")))
    n += 1
print(f"Invoice in advance: {n}")

conn.commit()

# === THỐNG KÊ ===
print("\n=== THỐNG KÊ ===")
for t in ("dm_vattu","dm_khachhang","phieu_giaodich","hop_dong","bang_gia"):
    print(f"  {t}: {conn.execute(f'SELECT COUNT(*) FROM {t}').fetchone()[0]}")
ton = conn.execute("SELECT SUM(so_luong) FROM phieu_giaodich WHERE loai_phieu='NHAP'").fetchone()[0] or 0
xuat = conn.execute("SELECT SUM(so_luong) FROM phieu_giaodich WHERE loai_phieu='XUAT'").fetchone()[0] or 0
kg = conn.execute("SELECT SUM(so_luong) FROM phieu_giaodich WHERE loai_phieu='KYGUI'").fetchone()[0] or 0
xkg = conn.execute("SELECT SUM(so_luong) FROM phieu_giaodich WHERE loai_phieu='XUAT_KYGUI'").fetchone()[0] or 0
print(f"  Tổng nhập: {ton} | Tổng xuất: {xuat} | Ký gửi: {kg} | Xuất ký gửi: {xkg}")
print(f"  Tồn kho (nhập - xuất): {ton - xuat}")
print(f"  Tồn ký gửi (kg - xkg): {kg - xkg}")
conn.close()
wb.close()
print("\nHoàn thành!")
